from __future__ import annotations

import hashlib
import json
import re
import shutil
import sqlite3
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Callable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

from openpyxl import load_workbook

from . import feishu_app_bot
from .professional_skills import ProfessionalSkillError, ProfessionalSkillRegistry


AUTHORIZED_TEST_GROUP_NAME = "智算测试"
MAX_OUTBOUND_RECIPIENTS = 10
REVIEW_COMMENT_MAX_LENGTH = 500
SOURCE_SYSTEM = "模拟造价系统"
EVENT_TYPE = "task.assigned"
DELIVERY_MODES = {"group", "direct", "mixed"}
DELIVERY_CHANNELS = {"group", "direct"}
DELIVERY_STAGES = ("task_card", "task_file", "review_card", "completion_card")
TASK_KIND = "external_dispatch"
SUBMISSION_WINDOW_SECONDS = 60
PUBLIC_STATUS_LABELS = {
    "pending_dispatch": "待派发",
    "delivering": "投递中",
    "pending_claim": "待领取",
    "claimed": "已领取",
    "awaiting_review_file": "待上传编制成果",
    "pending_review": "多人复核中",
    "returned": "已退回编制",
    "pending_upload": "待收件",
    "processing": "处理中",
    "completed": "已完成",
    "dispatch_failed": "投递失败",
    "failed": "失败",
}
TEXT_LIMITS = {
    "event_id": 120,
    "source_task_id": 120,
    "task_name": 200,
    "project_name": 160,
    "instructions": 2000,
    "template_asset_id": 200,
    "template_version": 80,
}


def generate_dispatch_source_task_id() -> str:
    return f"SIM-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex[:4].upper()}"


def generate_dispatch_project_name() -> str:
    return f"项目-{uuid4().hex[:6].upper()}"


def sanitize_dispatch_error(error: object) -> str:
    text = feishu_app_bot.sanitize_error(error)
    return re.sub(
        r"(?i)\b(app[_ -]?secret|access[_ -]?token|refresh[_ -]?token|token|password)\s*[:=]\s*[^\s,;]+",
        r"\1=***",
        text,
    )


class DispatchValidationError(ValueError):
    def __init__(self, message: str, *, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


def sanitize_review_comment(value: object) -> str:
    comment = str(value or "").replace("\x00", "").strip()
    if len(comment) > REVIEW_COMMENT_MAX_LENGTH:
        raise DispatchValidationError(f"复核评论不能超过 {REVIEW_COMMENT_MAX_LENGTH} 字")
    return comment


def display_review_comment(value: object) -> str:
    return sanitize_review_comment(value).replace("\r\n", "；").replace("\r", "；").replace("\n", "；")


def normalize_delivery_policy(value: object, *, legacy_mode: str = "group") -> dict[str, list[str]]:
    fallback = legacy_mode if legacy_mode in {"group", "direct"} else "group"
    raw = value if isinstance(value, dict) else {}
    normalized: dict[str, list[str]] = {}
    for stage in DELIVERY_STAGES:
        stage_value = raw.get(stage) if isinstance(raw, dict) else None
        if isinstance(stage_value, dict):
            channels = [channel for channel in ("group", "direct") if bool(stage_value.get(channel))]
        elif isinstance(stage_value, (list, tuple, set)):
            channels = [str(channel).strip() for channel in stage_value if str(channel).strip()]
        else:
            channels = [fallback]
        channels = list(dict.fromkeys(channels))
        if not channels or any(channel not in DELIVERY_CHANNELS for channel in channels):
            raise DispatchValidationError(f"{stage} 至少选择一个有效投递通道")
        normalized[stage] = channels
    return normalized


def delivery_mode_from_policy(policy: dict[str, list[str]]) -> str:
    channels = {channel for values in policy.values() for channel in values}
    if channels == {"group"}:
        return "group"
    if channels == {"direct"}:
        return "direct"
    return "mixed"


@dataclass(frozen=True)
class TaskArtifact:
    template_asset_id: str
    template_version: str


@dataclass(frozen=True)
class TaskEnvelope:
    event_id: str
    event_type: str
    source_system: str
    source_task_id: str
    task_name: str
    project_name: str
    skill_id: str
    skill_version: str
    delivery_mode: str
    platform_profile_id: str
    assignee_ref: str
    deadline: str
    instructions: str
    input_artifact: TaskArtifact
    reviewer_refs: tuple[str, ...] = ()
    delivery_policy: dict[str, list[str]] | None = None

    def validate(self) -> None:
        required = {
            "event_id": self.event_id,
            "source_task_id": self.source_task_id,
            "task_name": self.task_name,
            "project_name": self.project_name,
            "skill_id": self.skill_id,
            "platform_profile_id": self.platform_profile_id,
            "assignee_ref": self.assignee_ref,
            "deadline": self.deadline,
            "instructions": self.instructions,
            "template_asset_id": self.input_artifact.template_asset_id,
            "template_version": self.input_artifact.template_version,
        }
        missing = [name for name, value in required.items() if not str(value or "").strip()]
        if missing:
            raise DispatchValidationError(f"缺少必填字段：{', '.join(missing)}")
        if self.source_system != SOURCE_SYSTEM:
            raise DispatchValidationError("当前仅允许模拟造价系统发起 P0 任务")
        if self.event_type != EVENT_TYPE:
            raise DispatchValidationError("不支持的外部任务事件类型")
        if self.delivery_mode not in DELIVERY_MODES:
            raise DispatchValidationError("投递方式必须是 group、direct 或 mixed")
        self.normalized_delivery_policy()
        if not self.reviewer_refs:
            raise DispatchValidationError("至少选择一名复核人")
        if len(set(self.reviewer_refs)) != len(self.reviewer_refs):
            raise DispatchValidationError("复核人不能重复")
        named_people = {self.assignee_ref, *self.reviewer_refs}
        if len(named_people) > MAX_OUTBOUND_RECIPIENTS:
            raise DispatchValidationError(
                f"单个任务最多只能向 {MAX_OUTBOUND_RECIPIENTS} 名明确指定人员发送消息"
            )
        for name, value in required.items():
            limit = TEXT_LIMITS.get(name)
            if limit and len(str(value)) > limit:
                raise DispatchValidationError(f"{name} 超过 {limit} 字符上限")
        try:
            datetime.fromisoformat(self.deadline.replace("Z", "+00:00"))
        except ValueError as exc:
            raise DispatchValidationError("截止时间格式无效") from exc

    @property
    def business_key(self) -> str:
        return f"{self.source_system}\n{self.source_task_id}\n{self.event_type}"

    def normalized_delivery_policy(self) -> dict[str, list[str]]:
        return normalize_delivery_policy(self.delivery_policy, legacy_mode=self.delivery_mode)


class ExternalDispatchStore:
    def __init__(self, db_path: Path = feishu_app_bot.DB_PATH) -> None:
        self.db_path = Path(db_path)
        feishu_app_bot.TaskStore(self.db_path)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path, timeout=15)
        connection.row_factory = sqlite3.Row
        return connection

    def _init_db(self) -> None:
        columns: dict[str, str] = {
            "task_kind": "TEXT NOT NULL DEFAULT 'inbound_file'",
            "event_type": "TEXT NOT NULL DEFAULT ''",
            "source_system": "TEXT NOT NULL DEFAULT ''",
            "source_task_id": "TEXT NOT NULL DEFAULT ''",
            "business_key": "TEXT NOT NULL DEFAULT ''",
            "task_name": "TEXT NOT NULL DEFAULT ''",
            "project_name": "TEXT NOT NULL DEFAULT ''",
            "skill_id": "TEXT NOT NULL DEFAULT ''",
            "skill_version": "TEXT NOT NULL DEFAULT ''",
            "skill_snapshot_json": "TEXT NOT NULL DEFAULT '{}'",
            "delivery_mode": "TEXT NOT NULL DEFAULT ''",
            "delivery_policy_json": "TEXT NOT NULL DEFAULT '{}'",
            "delivery_state_json": "TEXT NOT NULL DEFAULT '{}'",
            "platform_profile_id": "TEXT NOT NULL DEFAULT ''",
            "target_chat_id": "TEXT NOT NULL DEFAULT ''",
            "target_chat_name": "TEXT NOT NULL DEFAULT ''",
            "assignee_mapping_id": "TEXT NOT NULL DEFAULT ''",
            "assignee_user_id": "TEXT NOT NULL DEFAULT ''",
            "assignee_name": "TEXT NOT NULL DEFAULT ''",
            "deadline": "TEXT NOT NULL DEFAULT ''",
            "instructions": "TEXT NOT NULL DEFAULT ''",
            "template_asset_id": "TEXT NOT NULL DEFAULT ''",
            "template_version": "TEXT NOT NULL DEFAULT ''",
            "template_hash": "TEXT NOT NULL DEFAULT ''",
            "template_source_path": "TEXT NOT NULL DEFAULT ''",
            "task_excel_path": "TEXT NOT NULL DEFAULT ''",
            "submission_file_name": "TEXT NOT NULL DEFAULT ''",
            "submission_excel_path": "TEXT NOT NULL DEFAULT ''",
            "submission_hash": "TEXT NOT NULL DEFAULT ''",
            "submission_message_id": "TEXT NOT NULL DEFAULT ''",
            "submission_received_at": "TEXT NOT NULL DEFAULT ''",
            "submission_window_expires_at": "TEXT NOT NULL DEFAULT ''",
            "submission_delivery_status": "TEXT NOT NULL DEFAULT ''",
            "submission_delivery_message_ids": "TEXT NOT NULL DEFAULT ''",
            "submission_error": "TEXT NOT NULL DEFAULT ''",
            "card_status": "TEXT NOT NULL DEFAULT 'pending'",
            "file_status": "TEXT NOT NULL DEFAULT 'pending'",
            "card_message_id": "TEXT NOT NULL DEFAULT ''",
            "file_message_id": "TEXT NOT NULL DEFAULT ''",
            "delivery_retry_count": "INTEGER NOT NULL DEFAULT 0",
            "delivery_error": "TEXT NOT NULL DEFAULT ''",
            "delivered_at": "TEXT NOT NULL DEFAULT ''",
            "claimed_at": "TEXT NOT NULL DEFAULT ''",
            "review_round": "INTEGER NOT NULL DEFAULT 0",
            "review_card_status": "TEXT NOT NULL DEFAULT ''",
            "review_card_message_id": "TEXT NOT NULL DEFAULT ''",
            "review_error": "TEXT NOT NULL DEFAULT ''",
            "completion_card_status": "TEXT NOT NULL DEFAULT ''",
            "completion_card_message_id": "TEXT NOT NULL DEFAULT ''",
            "completion_error": "TEXT NOT NULL DEFAULT ''",
            "completed_at": "TEXT NOT NULL DEFAULT ''",
        }
        with self._connect() as connection:
            existing = {str(row["name"]) for row in connection.execute("PRAGMA table_info(tasks)")}
            for name, definition in columns.items():
                if name not in existing:
                    connection.execute(f"ALTER TABLE tasks ADD COLUMN {name} {definition}")
            connection.executescript(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS idx_tasks_external_business_key
                ON tasks(business_key) WHERE task_kind='external_dispatch' AND business_key<>'';
                CREATE TABLE IF NOT EXISTS dispatch_personnel_mappings (
                    mapping_id TEXT PRIMARY KEY,
                    employee_id TEXT NOT NULL DEFAULT '',
                    account TEXT NOT NULL DEFAULT '',
                    display_name TEXT NOT NULL,
                    platform TEXT NOT NULL,
                    platform_user_id TEXT NOT NULL,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(platform, platform_user_id)
                );
                CREATE TABLE IF NOT EXISTS dispatch_attempts (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    task_id TEXT NOT NULL,
                    step TEXT NOT NULL,
                    status TEXT NOT NULL,
                    error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS dispatch_reviewers (
                    task_id TEXT NOT NULL,
                    mapping_id TEXT NOT NULL,
                    platform_user_id TEXT NOT NULL,
                    display_name TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'waiting',
                    review_round INTEGER NOT NULL DEFAULT 0,
                    comment TEXT NOT NULL DEFAULT '',
                    commented_at TEXT NOT NULL DEFAULT '',
                    decided_at TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(task_id, mapping_id)
                );
                CREATE TABLE IF NOT EXISTS dispatch_review_actions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    task_id TEXT NOT NULL,
                    review_round INTEGER NOT NULL,
                    reviewer_mapping_id TEXT NOT NULL,
                    reviewer_name TEXT NOT NULL,
                    decision TEXT NOT NULL,
                    comment TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL
                );
                """
            )
            reviewer_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(dispatch_reviewers)")}
            if "comment" not in reviewer_columns:
                connection.execute("ALTER TABLE dispatch_reviewers ADD COLUMN comment TEXT NOT NULL DEFAULT ''")
            if "commented_at" not in reviewer_columns:
                connection.execute("ALTER TABLE dispatch_reviewers ADD COLUMN commented_at TEXT NOT NULL DEFAULT ''")
            action_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(dispatch_review_actions)")}
            if "comment" not in action_columns:
                connection.execute("ALTER TABLE dispatch_review_actions ADD COLUMN comment TEXT NOT NULL DEFAULT ''")

    def known_chat_ids(self) -> list[str]:
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT DISTINCT chat_id FROM tasks WHERE chat_id<>'' UNION SELECT DISTINCT target_chat_id FROM tasks WHERE target_chat_id<>''"
            ).fetchall()
        return [str(row[0]) for row in rows if str(row[0] or "").strip()]

    def upsert_person(self, *, platform: str, platform_user_id: str, display_name: str) -> dict[str, Any]:
        mapping_id = "PM-" + hashlib.sha256(f"{platform}\n{platform_user_id}".encode("utf-8")).hexdigest()[:16].upper()
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                """INSERT INTO dispatch_personnel_mappings
                (mapping_id,display_name,platform,platform_user_id,enabled,created_at,updated_at)
                VALUES (?,?,?,?,1,?,?)
                ON CONFLICT(platform,platform_user_id) DO UPDATE SET
                display_name=excluded.display_name,updated_at=excluded.updated_at""",
                (mapping_id, display_name, platform, platform_user_id, now, now),
            )
            row = connection.execute(
                "SELECT * FROM dispatch_personnel_mappings WHERE platform=? AND platform_user_id=?",
                (platform, platform_user_id),
            ).fetchone()
        return dict(row)

    def get_person(self, mapping_id: str, platform: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM dispatch_personnel_mappings WHERE mapping_id=? AND platform=? AND enabled=1",
                (mapping_id, platform),
            ).fetchone()
        return dict(row) if row else None

    def find_business_task(self, business_key: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_kind=? AND business_key=?",
                (TASK_KIND, business_key),
            ).fetchone()
        return self._with_reviewers(dict(row)) if row else None

    def create_task(self, values: dict[str, Any]) -> tuple[dict[str, Any], bool]:
        columns = list(values)
        placeholders = ",".join("?" for _ in columns)
        with self._connect() as connection:
            try:
                connection.execute(
                    f"INSERT INTO tasks ({','.join(columns)}) VALUES ({placeholders})",
                    tuple(values[name] for name in columns),
                )
                created = True
            except sqlite3.IntegrityError:
                created = False
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_kind=? AND business_key=?",
                (TASK_KIND, values["business_key"]),
            ).fetchone()
        if not row:
            raise RuntimeError("外部任务创建失败")
        return dict(row), created

    def set_reviewers(self, task_id: str, reviewers: list[dict[str, Any]]) -> None:
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute("DELETE FROM dispatch_reviewers WHERE task_id=?", (task_id,))
            connection.executemany(
                """INSERT INTO dispatch_reviewers
                (task_id,mapping_id,platform_user_id,display_name,status,review_round,updated_at)
                VALUES (?,?,?,?, 'waiting',0,?)""",
                [(task_id, item["mapping_id"], item["platform_user_id"], item["display_name"], now) for item in reviewers],
            )

    def list_reviewers(self, task_id: str) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT * FROM dispatch_reviewers WHERE task_id=? ORDER BY display_name,mapping_id", (task_id,)
            ).fetchall()
        return [dict(row) for row in rows]

    def _with_reviewers(self, task: dict[str, Any]) -> dict[str, Any]:
        task["_reviewers"] = self.list_reviewers(str(task.get("task_id") or ""))
        return task

    def update_delivery(self, task_id: str, **fields: Any) -> dict[str, Any]:
        allowed = {
            "status", "stage", "error", "target_chat_id", "target_chat_name",
            "card_status", "file_status", "card_message_id", "file_message_id",
            "delivery_retry_count", "delivery_error", "delivered_at", "updated_at",
            "delivery_state_json",
        }
        updates = {key: value for key, value in fields.items() if key in allowed}
        updates["updated_at"] = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                f"UPDATE tasks SET {','.join(f'{name}=?' for name in updates)} WHERE task_id=? AND task_kind=?",
                (*updates.values(), task_id, TASK_KIND),
            )
            row = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        if not row:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        return self._with_reviewers(dict(row))

    def claim_task(self, task_id: str, *, operator_open_id: str, platform_profile_id: str) -> tuple[dict[str, Any], bool]:
        operator_id = str(operator_open_id or "").strip()
        profile_id = str(platform_profile_id or "").strip()
        if not operator_id:
            raise DispatchValidationError("无法确认领取人身份，请稍后重试", status_code=403)
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_id=? AND task_kind=?", (task_id, TASK_KIND),
            ).fetchone()
            if not row:
                raise DispatchValidationError("未找到外部派发任务", status_code=404)
            task = dict(row)
            if str(task.get("platform_profile_id") or "") != profile_id:
                raise DispatchValidationError("当前机器人不能领取其他平台的任务", status_code=403)
            if str(task.get("assignee_user_id") or "") != operator_id:
                raise DispatchValidationError("该任务已指定给其他编制人，您不能领取", status_code=403)
            if task.get("status") == "claimed":
                return self._with_reviewers(task), False
            if task.get("status") != "pending_claim" or task.get("file_status") != "sent":
                raise DispatchValidationError("任务尚未完成投递，暂时不能领取", status_code=409)
            connection.execute(
                "UPDATE tasks SET status='claimed',stage='claimed',claimed_at=?,updated_at=? WHERE task_id=? AND task_kind=?",
                (now, now, task_id, TASK_KIND),
            )
            claimed = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        return self._with_reviewers(dict(claimed)), True

    def open_submission_window(
        self,
        task_id: str,
        *,
        operator_open_id: str,
        platform_profile_id: str,
        now: datetime | None = None,
    ) -> tuple[dict[str, Any], bool]:
        operator_id = str(operator_open_id or "").strip()
        current = now or datetime.now(timezone.utc)
        if current.tzinfo is None:
            current = current.replace(tzinfo=timezone.utc)
        expires_at = (current.astimezone(timezone.utc) + timedelta(seconds=SUBMISSION_WINDOW_SECONDS)).isoformat(
            timespec="seconds"
        )
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_id=? AND task_kind=?",
                (task_id, TASK_KIND),
            ).fetchone()
            if not row:
                raise DispatchValidationError("未找到外部派发任务", status_code=404)
            task = dict(row)
            if str(task.get("platform_profile_id") or "") != str(platform_profile_id or ""):
                raise DispatchValidationError("当前机器人不能操作其他平台的任务", status_code=403)
            if str(task.get("assignee_user_id") or "") != operator_id:
                raise DispatchValidationError("只有目标编制人可以提交复核", status_code=403)
            if task.get("status") == "awaiting_review_file":
                try:
                    existing_expiry = datetime.fromisoformat(str(task.get("submission_window_expires_at") or ""))
                    if existing_expiry.tzinfo is None:
                        existing_expiry = existing_expiry.replace(tzinfo=timezone.utc)
                    if existing_expiry >= current.astimezone(timezone.utc):
                        return self._with_reviewers(task), False
                except ValueError:
                    pass
            if task.get("status") not in {"claimed", "returned", "awaiting_review_file"}:
                raise DispatchValidationError("当前任务尚不能提交多人复核", status_code=409)
            connection.execute(
                """UPDATE tasks SET status='awaiting_review_file',stage='awaiting_review_file',
                submission_window_expires_at=?,submission_delivery_status='',
                submission_delivery_message_ids='',submission_error='',updated_at=?
                WHERE task_id=? AND task_kind=?""",
                (expires_at, feishu_app_bot.utc_now(), task_id, TASK_KIND),
            )
            updated = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        return self._with_reviewers(dict(updated)), True

    def find_submission_window(
        self,
        *,
        operator_open_id: str,
        platform_profile_id: str,
        chat_id: str,
        is_private: bool,
        sent_at: str = "",
    ) -> dict[str, Any] | None:
        operator_id = str(operator_open_id or "").strip()
        if not operator_id:
            return None
        try:
            comparison_time = datetime.fromisoformat(str(sent_at or ""))
            if comparison_time.tzinfo is None:
                comparison_time = comparison_time.replace(tzinfo=timezone.utc)
        except ValueError:
            comparison_time = datetime.now(timezone.utc)
        with self._connect() as connection:
            rows = connection.execute(
                """SELECT * FROM tasks
                WHERE task_kind=? AND platform_profile_id=? AND assignee_user_id=?
                AND status='awaiting_review_file'
                ORDER BY updated_at DESC""",
                (TASK_KIND, str(platform_profile_id or ""), operator_id),
            ).fetchall()
        for row in rows:
            task = dict(row)
            try:
                expires = datetime.fromisoformat(str(task.get("submission_window_expires_at") or ""))
                if expires.tzinfo is None:
                    expires = expires.replace(tzinfo=timezone.utc)
            except ValueError:
                continue
            if comparison_time.astimezone(timezone.utc) > expires.astimezone(timezone.utc):
                continue
            if not is_private and str(task.get("target_chat_id") or "") != str(chat_id or ""):
                continue
            return self._with_reviewers(task)
        return None

    def submission_target_path(self, task: dict[str, Any], file_name: str) -> Path:
        task_id = str(task.get("task_id") or "").strip()
        if not task_id:
            raise DispatchValidationError("任务缺少内部编号")
        safe_name = feishu_app_bot.safe_filename(file_name)
        round_number = int(task.get("review_round") or 0) + 1
        return feishu_app_bot.RUNTIME_ROOT / "external-dispatch" / task_id / "submissions" / f"round-{round_number}" / safe_name

    def attach_submission(
        self,
        task_id: str,
        *,
        operator_open_id: str,
        platform_profile_id: str,
        message_id: str,
        file_name: str,
        file_path: Path,
    ) -> dict[str, Any]:
        path = Path(file_path)
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            raise DispatchValidationError("编制成果必须是可读取的 .xlsx 文件")
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_id=? AND task_kind=?",
                (task_id, TASK_KIND),
            ).fetchone()
            if not row:
                raise DispatchValidationError("未找到外部派发任务", status_code=404)
            task = dict(row)
            if str(task.get("platform_profile_id") or "") != str(platform_profile_id or ""):
                raise DispatchValidationError("当前机器人不能操作其他平台的任务", status_code=403)
            if str(task.get("assignee_user_id") or "") != str(operator_open_id or ""):
                raise DispatchValidationError("只有目标编制人可以提交成果", status_code=403)
            if task.get("status") != "awaiting_review_file":
                if str(task.get("submission_message_id") or "") == str(message_id or ""):
                    return self._with_reviewers(task)
                raise DispatchValidationError("当前任务没有等待编制成果", status_code=409)
            connection.execute(
                """UPDATE tasks SET submission_file_name=?,submission_excel_path=?,
                submission_hash=?,submission_message_id=?,submission_received_at=?,
                submission_delivery_status='pending',submission_error='',updated_at=?
                WHERE task_id=? AND task_kind=?""",
                (
                    feishu_app_bot.safe_filename(file_name),
                    str(path),
                    file_hash,
                    str(message_id or ""),
                    now,
                    now,
                    task_id,
                    TASK_KIND,
                ),
            )
            updated = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        return self._with_reviewers(dict(updated))

    def mark_submission_delivery(
        self,
        task_id: str,
        *,
        status: str,
        message_ids: list[str] | None = None,
        error: object = "",
    ) -> dict[str, Any]:
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                """UPDATE tasks SET submission_delivery_status=?,
                submission_delivery_message_ids=?,submission_error=?,updated_at=?
                WHERE task_id=? AND task_kind=?""",
                (
                    str(status or ""),
                    json.dumps([item for item in (message_ids or []) if item], ensure_ascii=False),
                    sanitize_dispatch_error(error),
                    now,
                    task_id,
                    TASK_KIND,
                ),
            )
        task = self.get_task(task_id)
        if not task:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        return task

    def submit_for_review(self, task_id: str, *, operator_open_id: str, platform_profile_id: str) -> tuple[dict[str, Any], bool]:
        operator_id = str(operator_open_id or "").strip()
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute("SELECT * FROM tasks WHERE task_id=? AND task_kind=?", (task_id, TASK_KIND)).fetchone()
            if not row:
                raise DispatchValidationError("未找到外部派发任务", status_code=404)
            task = dict(row)
            if str(task.get("platform_profile_id") or "") != str(platform_profile_id or ""):
                raise DispatchValidationError("当前机器人不能操作其他平台的任务", status_code=403)
            if str(task.get("assignee_user_id") or "") != operator_id:
                raise DispatchValidationError("只有目标编制人可以提交复核", status_code=403)
            if task.get("status") == "pending_review":
                return self._with_reviewers(task), False
            if task.get("status") not in {"claimed", "returned", "awaiting_review_file"}:
                raise DispatchValidationError("当前任务尚不能提交多人复核", status_code=409)
            if not str(task.get("submission_excel_path") or "").strip():
                raise DispatchValidationError("请先提交编制完成的 Excel 成果", status_code=409)
            reviewer_count = connection.execute("SELECT COUNT(*) FROM dispatch_reviewers WHERE task_id=?", (task_id,)).fetchone()[0]
            if not reviewer_count:
                raise DispatchValidationError("任务未配置复核人", status_code=409)
            review_round = int(task.get("review_round") or 0) + 1
            connection.execute(
                "UPDATE dispatch_reviewers SET status='pending',review_round=?,comment='',commented_at='',decided_at='',updated_at=? WHERE task_id=?",
                (review_round, now, task_id),
            )
            connection.execute(
                """UPDATE tasks SET status='pending_review',stage='pending_review',review_round=?,
                review_card_status='pending',review_card_message_id='',review_error='',updated_at=?
                WHERE task_id=? AND task_kind=?""",
                (review_round, now, task_id, TASK_KIND),
            )
            updated = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        return self._with_reviewers(dict(updated)), True

    def mark_review_card(self, task_id: str, *, status: str, message_id: str = "", error: str = "") -> dict[str, Any]:
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                "UPDATE tasks SET review_card_status=?,review_card_message_id=?,review_error=?,updated_at=? WHERE task_id=? AND task_kind=?",
                (status, message_id, sanitize_dispatch_error(error), now, task_id, TASK_KIND),
            )
        task = self.get_task(task_id)
        if not task:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        return task

    def mark_completion_card(self, task_id: str, *, status: str, message_id: str = "", error: str = "") -> dict[str, Any]:
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                "UPDATE tasks SET completion_card_status=?,completion_card_message_id=?,completion_error=?,updated_at=? WHERE task_id=? AND task_kind=?",
                (status, message_id, sanitize_dispatch_error(error), now, task_id, TASK_KIND),
            )
        task = self.get_task(task_id)
        if not task:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        return task

    def rollback_review_submission(self, task_id: str, error: object) -> dict[str, Any]:
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute(
                "UPDATE tasks SET status='claimed',stage='review_card_failed',review_card_status='failed',review_error=?,updated_at=? WHERE task_id=? AND task_kind=?",
                (sanitize_dispatch_error(error), now, task_id, TASK_KIND),
            )
        return self.get_task(task_id) or {}

    def review_task(
        self,
        task_id: str,
        *,
        operator_open_id: str,
        platform_profile_id: str,
        decision: str,
        comment: str = "",
    ) -> tuple[dict[str, Any], bool]:
        if decision not in {"approve", "reject"}:
            raise DispatchValidationError("不支持的复核结论")
        review_comment = sanitize_review_comment(comment)
        if decision == "reject" and not review_comment:
            raise DispatchValidationError("退回编制时必须填写复核评论")
        operator_id = str(operator_open_id or "").strip()
        now = feishu_app_bot.utc_now()
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute("SELECT * FROM tasks WHERE task_id=? AND task_kind=?", (task_id, TASK_KIND)).fetchone()
            if not row:
                raise DispatchValidationError("未找到外部派发任务", status_code=404)
            task = dict(row)
            if str(task.get("platform_profile_id") or "") != str(platform_profile_id or ""):
                raise DispatchValidationError("当前机器人不能操作其他平台的任务", status_code=403)
            review_round = int(task.get("review_round") or 0)
            reviewer = connection.execute(
                "SELECT * FROM dispatch_reviewers WHERE task_id=? AND platform_user_id=? AND review_round=?",
                (task_id, operator_id, review_round),
            ).fetchone()
            if not reviewer:
                raise DispatchValidationError("您不是本任务的复核人", status_code=403)
            reviewer_data = dict(reviewer)
            target_status = "approved" if decision == "approve" else "rejected"
            if task.get("status") != "pending_review":
                if reviewer_data.get("status") == target_status:
                    return self._with_reviewers(task), False
                raise DispatchValidationError("本轮复核已经结束", status_code=409)
            if reviewer_data.get("status") != "pending":
                if reviewer_data.get("status") == target_status:
                    return self._with_reviewers(task), False
                raise DispatchValidationError("您已提交本轮复核结论", status_code=409)
            connection.execute(
                """UPDATE dispatch_reviewers SET status=?,comment=?,commented_at=?,
                decided_at=?,updated_at=? WHERE task_id=? AND mapping_id=?""",
                (target_status, review_comment, now if review_comment else "", now, now, task_id, reviewer_data["mapping_id"]),
            )
            connection.execute(
                """INSERT INTO dispatch_review_actions
                (task_id,review_round,reviewer_mapping_id,reviewer_name,decision,comment,created_at)
                VALUES (?,?,?,?,?,?,?)""",
                (
                    task_id,
                    review_round,
                    reviewer_data["mapping_id"],
                    reviewer_data["display_name"],
                    decision,
                    review_comment,
                    now,
                ),
            )
            if decision == "reject":
                connection.execute(
                    "UPDATE tasks SET status='returned',stage='returned',updated_at=? WHERE task_id=?", (now, task_id)
                )
            else:
                pending = connection.execute(
                    "SELECT COUNT(*) FROM dispatch_reviewers WHERE task_id=? AND review_round=? AND status='pending'",
                    (task_id, review_round),
                ).fetchone()[0]
                if pending == 0:
                    connection.execute(
                        "UPDATE tasks SET status='completed',stage='completed',completed_at=?,updated_at=? WHERE task_id=?",
                        (now, now, task_id),
                    )
            updated = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        return self._with_reviewers(dict(updated)), True

    def record_attempt(self, task_id: str, step: str, status: str, error: str = "") -> None:
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO dispatch_attempts(task_id,step,status,error,created_at) VALUES (?,?,?,?,?)",
                (task_id, step, status, sanitize_dispatch_error(error), feishu_app_bot.utc_now()),
            )

    def get_task(self, task_id: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM tasks WHERE task_id=? AND task_kind=?", (task_id, TASK_KIND),
            ).fetchone()
        return self._with_reviewers(dict(row)) if row else None

    def list_tasks(self, limit: int = 30) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT * FROM tasks WHERE task_kind=? ORDER BY created_at DESC LIMIT ?",
                (TASK_KIND, max(1, min(int(limit), 100))),
            ).fetchall()
        return [self._with_reviewers(dict(row)) for row in rows]


class ExternalTaskDispatchService:
    def __init__(
        self,
        *,
        store: ExternalDispatchStore,
        registry: ProfessionalSkillRegistry,
        feishu: Any,
        profile_id: str,
        runtime_root: Path = feishu_app_bot.RUNTIME_ROOT,
        authorized_group_name: str = AUTHORIZED_TEST_GROUP_NAME,
        app_url: str = "",
        direct_delivery_verified: bool = False,
    ) -> None:
        self.store = store
        self.registry = registry
        self.feishu = feishu
        self.profile_id = profile_id
        self.runtime_root = Path(runtime_root)
        self.authorized_group_name = authorized_group_name
        self.app_url = app_url
        self.direct_delivery_verified = direct_delivery_verified

    @property
    def directory_cache_path(self) -> Path:
        return self.runtime_root / "external-dispatch" / "directory-cache.json"

    def _read_directory_cache(self) -> dict[str, Any] | None:
        try:
            payload = json.loads(self.directory_cache_path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return None
        profiles = payload.get("profiles") if isinstance(payload, dict) else None
        cached = profiles.get(self.profile_id) if isinstance(profiles, dict) else None
        return cached if isinstance(cached, dict) else None

    def _write_directory_cache(self, directory: dict[str, Any]) -> None:
        path = self.directory_cache_path
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            payload = {}
        if not isinstance(payload, dict):
            payload = {}
        profiles = payload.get("profiles") if isinstance(payload.get("profiles"), dict) else {}
        profiles[self.profile_id] = directory
        payload.update({"version": 1, "profiles": profiles})
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def refresh_directory(self) -> dict[str, Any]:
        previous = self._read_directory_cache() or {}
        previous_groups = {
            str(item.get("chat_id") or ""): item
            for item in previous.get("groups") or []
            if isinstance(item, dict) and str(item.get("chat_id") or "").strip()
        }
        groups: list[dict[str, Any]] = []
        for chat in self.feishu.list_chats():
            chat_id = str(chat.get("chat_id") or "").strip()
            if not chat_id:
                continue
            name = str(chat.get("name") or "").strip() or self.feishu.resolve_chat_name(chat_id) or "未命名群聊"
            members_available = True
            member_error = ""
            people: list[dict[str, str]] = []
            member_count = 0
            try:
                members_payload = self.feishu.list_chat_members(chat_id)
                member_count = int(members_payload.get("member_total") or 0)
                for member in members_payload.get("members") or []:
                    user_id = str(member.get("member_id") or "").strip()
                    display_name = str(member.get("name") or "").strip() or "未命名成员"
                    if not user_id:
                        continue
                    mapping = self.store.upsert_person(
                        platform=self.profile_id,
                        platform_user_id=user_id,
                        display_name=display_name,
                    )
                    people.append({"person_ref": str(mapping["mapping_id"]), "display_name": display_name})
                people = list({item["person_ref"]: item for item in people}.values())
                people.sort(key=lambda item: item["display_name"])
                member_count = max(member_count, len(people))
            except Exception as exc:
                members_available = False
                member_error = sanitize_dispatch_error(exc)
                cached_group = previous_groups.get(chat_id) or {}
                people = [item for item in cached_group.get("people") or [] if isinstance(item, dict)]
                member_count = int(cached_group.get("member_count") or len(people))
            groups.append({
                "group_ref": "GR-" + hashlib.sha256(f"{self.profile_id}\n{chat_id}".encode("utf-8")).hexdigest()[:16].upper(),
                "chat_id": chat_id,
                "name": str(name),
                "member_count": member_count,
                "members_available": members_available,
                "member_error": member_error,
                "people": people,
            })
        groups.sort(key=lambda item: (item["name"] != self.authorized_group_name, item["name"]))
        directory = {"updated_at": feishu_app_bot.utc_now(), "groups": groups}
        self._write_directory_cache(directory)
        return directory

    def directory(self, *, refresh: bool = False) -> dict[str, Any]:
        cached = self._read_directory_cache()
        refresh_error = ""
        if refresh or not cached:
            try:
                cached = self.refresh_directory()
                source = "live"
            except Exception as exc:
                if not cached:
                    raise
                source = "cache"
                refresh_error = sanitize_dispatch_error(exc)
        else:
            source = "cache"
        public_groups = [
            {
                "group_ref": str(group.get("group_ref") or ""),
                "name": str(group.get("name") or "未命名群聊"),
                "member_count": int(group.get("member_count") or 0),
                "members_available": bool(group.get("members_available", True)),
                "authorized": str(group.get("name") or "") == self.authorized_group_name,
                "people": [
                    {
                        "person_ref": str(person.get("person_ref") or ""),
                        "display_name": str(person.get("display_name") or "未命名成员"),
                    }
                    for person in group.get("people") or []
                    if isinstance(person, dict) and str(person.get("person_ref") or "").strip()
                ],
            }
            for group in cached.get("groups") or []
            if isinstance(group, dict) and str(group.get("group_ref") or "").strip()
        ]
        all_people = {
            person["person_ref"]: person
            for group in public_groups
            for person in group["people"]
        }
        return {
            "updated_at": str(cached.get("updated_at") or ""),
            "source": source,
            "refresh_error": refresh_error,
            "groups": public_groups,
            "people": sorted(all_people.values(), key=lambda item: item["display_name"]),
        }

    def options(self, *, refresh_directory: bool = False) -> dict[str, Any]:
        directory = self.directory(refresh=refresh_directory)
        authorized_groups = [item for item in directory["groups"] if item["authorized"]]
        if len(authorized_groups) != 1:
            reason = "未找到" if not authorized_groups else "找到多个同名"
            raise DispatchValidationError(
                f"{reason}唯一授权群“{self.authorized_group_name}”，已拒绝真实投递",
                status_code=409,
            )
        target_group = authorized_groups[0]
        people = list(target_group["people"])
        return {
            "source_system": SOURCE_SYSTEM,
            "event_type": EVENT_TYPE,
            "target_group": {
                "name": self.authorized_group_name,
                "available": target_group is not None,
            },
            "people": people,
            "directory": directory,
            "direct_delivery": {
                "status": "available" if self.direct_delivery_verified else "pending_verification",
                "label": "可用" if self.direct_delivery_verified else "待验证",
            },
        }

    def resolve_authorized_group(self) -> tuple[str, str]:
        candidates: dict[str, str] = {}
        cached = self._read_directory_cache() or {}
        for group in cached.get("groups") or []:
            if not isinstance(group, dict):
                continue
            chat_id = str(group.get("chat_id") or "").strip()
            if chat_id:
                candidates[chat_id] = str(group.get("name") or "").strip()
        for chat_id in self.store.known_chat_ids():
            try:
                candidates[chat_id] = str(self.feishu.resolve_chat_name(chat_id) or "").strip()
            except Exception:
                continue
        try:
            for chat in self.feishu.list_chats():
                chat_id = str(chat.get("chat_id") or "").strip()
                name = str(chat.get("name") or "").strip()
                if chat_id:
                    candidates[chat_id] = name
        except Exception:
            if not candidates:
                raise
        matches = [(chat_id, name) for chat_id, name in candidates.items() if name == self.authorized_group_name]
        if len(matches) != 1:
            reason = "未找到" if not matches else "找到多个同名"
            raise DispatchValidationError(
                f"{reason}唯一授权群“{self.authorized_group_name}”，已拒绝真实投递",
                status_code=409,
            )
        return matches[0]

    def create_and_deliver(self, envelope: TaskEnvelope, *, file_name: str, file_bytes: bytes) -> tuple[dict[str, Any], bool]:
        envelope.validate()
        delivery_policy = envelope.normalized_delivery_policy()
        if envelope.platform_profile_id != self.profile_id:
            raise DispatchValidationError("投递平台与当前机器人配置不一致", status_code=409)
        existing = self.store.find_business_task(envelope.business_key)
        if existing:
            return public_dispatch_task(existing), False
        try:
            skill_snapshot = self.registry.resolve_for_task(envelope.skill_id, envelope.skill_version or None)
        except ProfessionalSkillError as exc:
            raise DispatchValidationError(exc.message, status_code=exc.status_code) from exc
        person = self.store.get_person(envelope.assignee_ref, envelope.platform_profile_id)
        if not person:
            raise DispatchValidationError("目标编制人未建立可用的平台人员映射", status_code=409)
        if not str(person.get("display_name") or "").strip():
            raise DispatchValidationError("目标编制人缺少明确姓名，已拒绝投递", status_code=409)
        reviewers: list[dict[str, Any]] = []
        for reviewer_ref in envelope.reviewer_refs:
            reviewer = self.store.get_person(reviewer_ref, envelope.platform_profile_id)
            if not reviewer:
                raise DispatchValidationError("存在未建立平台人员映射的复核人", status_code=409)
            if not str(reviewer.get("display_name") or "").strip():
                raise DispatchValidationError("存在未明确姓名的复核人，已拒绝投递", status_code=409)
            reviewers.append(reviewer)
        if any("direct" in channels for channels in delivery_policy.values()) and not self.direct_delivery_verified:
            raise DispatchValidationError("当前机器人的主动单聊触达能力尚未完成验证", status_code=409)
        self._validate_xlsx(file_name, file_bytes)

        task_id = f"FS-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex[:6].upper()}"
        template_hash = hashlib.sha256(file_bytes).hexdigest()
        template_dir = self.runtime_root / "external-dispatch" / "templates"
        task_dir = self.runtime_root / "tasks" / task_id / "dispatch"
        template_dir.mkdir(parents=True, exist_ok=True)
        task_dir.mkdir(parents=True, exist_ok=True)
        source_path = template_dir / f"{template_hash}.xlsx"
        if not source_path.exists():
            source_path.write_bytes(file_bytes)
        task_file_name = self._task_file_name(task_id, envelope.project_name)
        task_path = task_dir / task_file_name
        shutil.copy2(source_path, task_path)
        now = feishu_app_bot.utc_now()
        values = {
            "task_id": task_id,
            "event_id": envelope.event_id,
            "message_id": f"dispatch:{task_id}",
            "chat_id": "",
            "file_key": f"dispatch:{task_id}:{template_hash[:12]}",
            "file_name": task_file_name,
            "status": "pending_dispatch",
            "stage": "pending_dispatch",
            "created_at": now,
            "updated_at": now,
            "task_kind": TASK_KIND,
            "event_type": envelope.event_type,
            "source_system": envelope.source_system,
            "source_task_id": envelope.source_task_id,
            "business_key": envelope.business_key,
            "task_name": envelope.task_name,
            "project_name": envelope.project_name,
            "skill_id": str(skill_snapshot.get("id") or envelope.skill_id),
            "skill_version": str(skill_snapshot.get("version") or envelope.skill_version),
            "skill_snapshot_json": json.dumps(skill_snapshot, ensure_ascii=False, separators=(",", ":")),
            "delivery_mode": delivery_mode_from_policy(delivery_policy),
            "delivery_policy_json": json.dumps(delivery_policy, ensure_ascii=False, separators=(",", ":")),
            "delivery_state_json": "{}",
            "platform_profile_id": envelope.platform_profile_id,
            "assignee_mapping_id": str(person["mapping_id"]),
            "assignee_user_id": str(person["platform_user_id"]),
            "assignee_name": str(person["display_name"]),
            "deadline": envelope.deadline,
            "instructions": envelope.instructions,
            "template_asset_id": envelope.input_artifact.template_asset_id,
            "template_version": envelope.input_artifact.template_version,
            "template_hash": template_hash,
            "template_source_path": self._runtime_relative(source_path),
            "task_excel_path": self._runtime_relative(task_path),
        }
        task, created = self.store.create_task(values)
        if not created:
            shutil.rmtree(task_dir.parent, ignore_errors=True)
            return public_dispatch_task(task), False
        self.store.set_reviewers(task_id, reviewers)
        task = self.deliver(task_id)
        return public_dispatch_task(task), True

    def deliver(self, task_id: str, *, retry: bool = False) -> dict[str, Any]:
        task = self.store.get_task(task_id)
        if not task:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        if task["card_status"] == "sent" and task["file_status"] == "sent":
            return task
        if retry:
            task = self.store.update_delivery(
                task_id,
                delivery_retry_count=int(task.get("delivery_retry_count") or 0) + 1,
                delivery_error="",
            )
        try:
            if any("group" in channels for channels in self._policy_for_task(task).values()):
                target_chat_id, target_chat_name = self._ensure_group_target(task)
            else:
                target_chat_id, target_chat_name = "", "精准单聊"
            task = self.store.update_delivery(
                task_id,
                status="delivering",
                stage="delivering",
                error="",
                target_chat_id=target_chat_id,
                target_chat_name=target_chat_name,
            )
            if task["file_status"] != "sent":
                try:
                    path = self.runtime_root / str(task["task_excel_path"])
                    task, message_ids = self._deliver_stage(
                        task,
                        "task_file",
                        lambda receive_id, receive_id_type: self.feishu.send_file_to(receive_id, receive_id_type, path),
                    )
                    task = self.store.update_delivery(
                        task_id,
                        file_status="sent",
                        file_message_id=json.dumps(message_ids, ensure_ascii=False),
                    )
                except Exception as exc:
                    return self._delivery_failed(task_id, "file", exc)
            if task["card_status"] != "sent":
                try:
                    task = self.store.update_delivery(
                        task_id,
                        status="pending_claim",
                        stage="pending_claim",
                        delivery_error="",
                    )
                    task, message_ids = self._deliver_stage(
                        task,
                        "task_card",
                        lambda receive_id, receive_id_type: self.feishu.send_card_to(
                            receive_id,
                            receive_id_type,
                            build_external_task_card(task, app_url=self._task_url(task_id)),
                        ),
                    )
                    task = self.store.update_delivery(
                        task_id,
                        card_status="sent",
                        card_message_id=json.dumps(message_ids, ensure_ascii=False),
                    )
                except Exception as exc:
                    return self._delivery_failed(task_id, "card", exc)
            if task.get("status") != "pending_claim":
                return task
            return self.store.update_delivery(
                task_id,
                status="pending_claim",
                stage="pending_claim",
                delivery_error="",
                delivered_at=feishu_app_bot.utc_now(),
            )
        except Exception as exc:
            return self._delivery_failed(task_id, "target", exc)

    def retry(self, task_id: str) -> dict[str, Any]:
        task = self.store.get_task(task_id)
        if not task:
            raise DispatchValidationError("未找到外部派发任务", status_code=404)
        if task["status"] not in {"dispatch_failed", "pending_dispatch", "delivering"}:
            raise DispatchValidationError("当前任务没有可重试的投递步骤", status_code=409)
        return public_dispatch_task(self.deliver(task_id, retry=True))

    @staticmethod
    def _json_dict(value: object) -> dict[str, Any]:
        if isinstance(value, dict):
            return value
        try:
            parsed = json.loads(str(value or "{}"))
        except (TypeError, ValueError):
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _policy_for_task(self, task: dict[str, Any]) -> dict[str, list[str]]:
        return normalize_delivery_policy(
            self._json_dict(task.get("delivery_policy_json")),
            legacy_mode=str(task.get("delivery_mode") or "group"),
        )

    def _ensure_group_target(self, task: dict[str, Any]) -> tuple[str, str]:
        if task.get("target_chat_id") and task.get("target_chat_name") == self.authorized_group_name:
            return str(task["target_chat_id"]), str(task["target_chat_name"])
        return self.resolve_authorized_group()

    def _delivery_targets(self, task: dict[str, Any], stage: str) -> list[tuple[str, str]]:
        channels = self._policy_for_task(task)[stage]
        targets: list[tuple[str, str]] = []
        if "group" in channels:
            chat_id, _ = self._ensure_group_target(task)
            targets.append((chat_id, "chat_id"))
        if "direct" in channels:
            if not self.direct_delivery_verified:
                raise DispatchValidationError("主动单聊能力尚未验证", status_code=409)
            user_id = str(task.get("assignee_user_id") or "").strip()
            if not user_id:
                raise DispatchValidationError("目标编制人没有稳定平台 ID", status_code=409)
            targets.append((user_id, "open_id"))
        return list(dict.fromkeys(targets))

    def _deliver_stage(
        self,
        task: dict[str, Any],
        stage: str,
        sender: Callable[[str, str], str],
    ) -> tuple[dict[str, Any], list[str]]:
        state = self._json_dict(task.get("delivery_state_json"))
        stage_state = state.get(stage) if isinstance(state.get(stage), dict) else {}
        message_ids: list[str] = []
        targets = self._delivery_targets(task, stage)
        enforce_outbound_audience_safety(
            self.feishu,
            targets,
            named_recipients={
                str(task.get("assignee_user_id") or "").strip(): str(task.get("assignee_name") or "").strip(),
            },
        )
        for receive_id, receive_id_type in targets:
            target_key = f"{receive_id_type}:{hashlib.sha256(receive_id.encode('utf-8')).hexdigest()[:16]}"
            current = stage_state.get(target_key) if isinstance(stage_state.get(target_key), dict) else {}
            if current.get("status") == "sent":
                if current.get("message_id"):
                    message_ids.append(str(current["message_id"]))
                continue
            try:
                message_id = sender(receive_id, receive_id_type) or ""
                stage_state[target_key] = {"status": "sent", "message_id": message_id, "error": ""}
                self.store.record_attempt(task["task_id"], f"{stage}:{receive_id_type}", "sent")
                if message_id:
                    message_ids.append(message_id)
            except Exception as exc:
                stage_state[target_key] = {"status": "failed", "message_id": "", "error": sanitize_dispatch_error(exc)}
                state[stage] = stage_state
                task = self.store.update_delivery(
                    task["task_id"],
                    delivery_state_json=json.dumps(state, ensure_ascii=False, separators=(",", ":")),
                )
                self.store.record_attempt(task["task_id"], f"{stage}:{receive_id_type}", "failed", str(exc))
                raise
            state[stage] = stage_state
            task = self.store.update_delivery(
                task["task_id"],
                delivery_state_json=json.dumps(state, ensure_ascii=False, separators=(",", ":")),
            )
        return task, message_ids

    def _delivery_failed(self, task_id: str, step: str, exc: Exception) -> dict[str, Any]:
        error = sanitize_dispatch_error(exc)
        return self.store.update_delivery(
            task_id,
            status="dispatch_failed",
            stage=f"{step}_failed",
            error=error,
            delivery_error=error,
        )

    def _task_url(self, task_id: str) -> str:
        if not self.app_url:
            return ""
        parsed = urlsplit(self.app_url)
        query = dict(parse_qsl(parsed.query, keep_blank_values=True))
        query["external_task_id"] = task_id
        return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urlencode(query), parsed.fragment))

    def _runtime_relative(self, path: Path) -> str:
        return path.resolve().relative_to(self.runtime_root.resolve()).as_posix()

    @staticmethod
    def _task_file_name(task_id: str, project_name: str) -> str:
        project = re.sub(r"[\\/:*?\"<>|\r\n]+", "-", project_name).strip(" .-")[:60] or "项目"
        return f"{task_id}-{project}-待填写.xlsx"

    @staticmethod
    def _validate_xlsx(file_name: str, file_bytes: bytes) -> None:
        if not str(file_name or "").lower().endswith(".xlsx"):
            raise DispatchValidationError("待填模板仅允许 .xlsx 文件")
        if not file_bytes:
            raise DispatchValidationError("待填模板为空")
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
            workbook.close()
        except Exception as exc:
            raise DispatchValidationError("待填模板不是可读取的 .xlsx 文件") from exc


def build_external_task_card(task: dict[str, Any], *, app_url: str = "") -> dict[str, Any]:
    assignee_name = feishu_app_bot.safe_filename(str(task.get("assignee_name") or "未指定"))
    assignee_id = str(task.get("assignee_user_id") or "").strip()
    assignee_line = assignee_name
    if "group" in stored_delivery_policy(task)["task_card"] and assignee_id:
        assignee_line = f"<at id={assignee_id}></at>（{assignee_name}）"
    status = str(task.get("status") or "")
    status_text = PUBLIC_STATUS_LABELS.get(status, status)
    reviewers = task.get("_reviewers") or []
    reviewer_names = "、".join(str(item.get("display_name") or "未命名") for item in reviewers) or "未配置"
    submission_name = str(task.get("submission_file_name") or "").strip()
    submission_line = f"\n**编制成果：** {submission_name}" if submission_name else ""
    upload_hint = (
        "\n\n**下一步：** 请在 1 分钟内直接发送一个编制完成的 `.xlsx` 文件；"
        "机器人会将该文件绑定到本任务并自动发起多人复核。"
        if status == "awaiting_review_file"
        else ""
    )
    content = (
        f"**状态：** {status_text}\n"
        f"**任务名称：** {task.get('task_name') or '-'}\n"
        f"**内部任务编号：** {task.get('task_id') or '-'}\n"
        f"**外部任务编号：** {task.get('source_task_id') or '-'}\n"
        f"**来源系统：** {task.get('source_system') or '-'}\n"
        f"**目标编制人：** {assignee_line}\n"
        f"**复核人：** {reviewer_names}\n"
        f"**流程：** 编制 → 多人复核（{len(reviewers)}人）→ 完成\n"
        f"**截止时间：** {task.get('deadline') or '-'}\n"
        f"**专业能力：** {task.get('skill_id') or '-'} · {task.get('skill_version') or '-'}\n\n"
        f"**任务说明：** {task.get('instructions') or '-'}"
        f"{submission_line}{upload_hint}"
    )
    elements: list[dict[str, Any]] = [{"tag": "div", "text": {"tag": "lark_md", "content": content}}]
    actions: list[dict[str, Any]] = []
    if status in {"pending_dispatch", "delivering", "pending_claim"}:
        actions.append({
            "tag": "button",
            "text": {"tag": "plain_text", "content": "领取任务"},
            "type": "primary",
            "value": {"action": "claim_external_task", "task_id": str(task.get("task_id") or "")},
        })
    elif status in {"claimed", "returned"}:
        actions.append({
            "tag": "button",
            "text": {
                "tag": "plain_text",
                "content": "提交成果并进入复核" if status == "claimed" else "重新提交成果并进入复核",
            },
            "type": "primary",
            "value": {"action": "submit_external_review", "task_id": str(task.get("task_id") or "")},
        })
    if app_url:
        actions.append({
                "tag": "button",
                "text": {"tag": "plain_text", "content": "进入工作台"},
                "type": "default",
                "url": app_url,
        })
    if actions:
        elements.append({"tag": "action", "actions": actions})
    return {
        "config": {"wide_screen_mode": True},
        "header": {
            "title": {"tag": "plain_text", "content": "造价智算 · P0模拟派发新任务"},
            "template": "green" if status == "completed" else ("yellow" if status == "returned" else "blue"),
        },
        "elements": elements,
    }


def build_external_review_card(task: dict[str, Any]) -> dict[str, Any]:
    status = str(task.get("status") or "")
    labels = {"waiting": "待编制", "pending": "待复核", "approved": "已通过", "rejected": "已退回"}
    reviewer_lines = "\n".join(
        (
            f"- {item.get('display_name') or '未命名'}："
            f"{labels.get(str(item.get('status') or ''), item.get('status') or '-')}"
            + (
                f"\n  评论：{display_review_comment(item.get('comment'))}"
                if str(item.get("comment") or "").strip()
                else ""
            )
        )
        for item in (task.get("_reviewers") or [])
    ) or "- 未配置复核人"
    content = (
        f"**状态：** {PUBLIC_STATUS_LABELS.get(status, status)}\n"
        f"**任务：** {task.get('task_name') or '-'}\n"
        f"**项目：** {task.get('project_name') or '-'}\n"
        f"**任务编号：** {task.get('task_id') or '-'}\n"
        f"**编制人：** {task.get('assignee_name') or '-'}\n"
        f"**编制成果：** {task.get('submission_file_name') or '未绑定'}\n"
        f"**复核轮次：** 第 {int(task.get('review_round') or 0)} 轮\n\n"
        f"**复核人员：**\n{reviewer_lines}"
    )
    elements: list[dict[str, Any]] = [{"tag": "div", "text": {"tag": "lark_md", "content": content}}]
    if status == "pending_review":
        elements.append({
            "tag": "form",
            "name": "review_form",
            "elements": [
                {
                    "tag": "input",
                    "name": "review_comment",
                    "label": {"tag": "plain_text", "content": "复核评论"},
                    "label_position": "top",
                    "placeholder": {"tag": "plain_text", "content": "通过时可选；退回编制时必填，最多 500 字"},
                    "max_length": REVIEW_COMMENT_MAX_LENGTH,
                },
                {
                    "tag": "button",
                    "name": "approve_review",
                    "action_type": "form_submit",
                    "text": {"tag": "plain_text", "content": "复核通过"},
                    "type": "primary",
                    "value": {
                        "action": "review_external_task",
                        "decision": "approve",
                        "task_id": str(task.get("task_id") or ""),
                    },
                },
                {
                    "tag": "button",
                    "name": "reject_review",
                    "action_type": "form_submit",
                    "text": {"tag": "plain_text", "content": "退回编制"},
                    "type": "danger",
                    "value": {
                        "action": "review_external_task",
                        "decision": "reject",
                        "task_id": str(task.get("task_id") or ""),
                    },
                },
            ],
        })
    return {
        "config": {"wide_screen_mode": True},
        "header": {"title": {"tag": "plain_text", "content": "造价智算 · 多人复核"}, "template": "green" if status == "completed" else ("yellow" if status == "returned" else "blue")},
        "elements": elements,
    }


def review_delivery_targets(
    task: dict[str, Any],
    *,
    authorized_group_name: str = AUTHORIZED_TEST_GROUP_NAME,
) -> list[tuple[str, str]]:
    """Resolve review-card recipients from the task's selected delivery policy."""
    channels = stored_delivery_policy(task)["review_card"]
    targets: list[tuple[str, str]] = []
    if "direct" in channels:
        recipients = _explicit_named_recipients(task.get("_reviewers") or [], role="复核人")
        if not recipients:
            raise DispatchValidationError("复核人没有稳定平台 ID，无法个人投递", status_code=409)
        targets.extend((user_id, "open_id") for user_id in sorted(recipients))

    if "group" in channels:
        chat_id = str(task.get("target_chat_id") or "").strip()
        chat_name = str(task.get("target_chat_name") or "").strip()
        if not chat_id or chat_name != authorized_group_name:
            raise DispatchValidationError("复核卡片仅允许投递到唯一授权群", status_code=409)
        targets.append((chat_id, "chat_id"))
    return list(dict.fromkeys(targets))


def completion_delivery_targets(
    task: dict[str, Any],
    *,
    authorized_group_name: str = AUTHORIZED_TEST_GROUP_NAME,
) -> list[tuple[str, str]]:
    channels = stored_delivery_policy(task)["completion_card"]
    targets: list[tuple[str, str]] = []
    if "direct" in channels:
        recipients = _explicit_named_recipients(
            [
                {
                    "platform_user_id": task.get("assignee_user_id"),
                    "display_name": task.get("assignee_name"),
                },
                *(task.get("_reviewers") or []),
            ],
            role="完结通知接收人",
        )
        if not recipients:
            raise DispatchValidationError("完结通知没有可用的个人接收人", status_code=409)
        targets.extend((user_id, "open_id") for user_id in sorted(recipients))
    if "group" in channels:
        chat_id = str(task.get("target_chat_id") or "").strip()
        chat_name = str(task.get("target_chat_name") or "").strip()
        if not chat_id or chat_name != authorized_group_name:
            raise DispatchValidationError("完结通知仅允许投递到唯一授权群", status_code=409)
        targets.append((chat_id, "chat_id"))
    return list(dict.fromkeys(targets))


def _explicit_named_recipients(people: list[dict[str, Any]], *, role: str) -> dict[str, str]:
    recipients: dict[str, str] = {}
    for person in people:
        user_id = str(person.get("platform_user_id") or "").strip()
        name = str(person.get("display_name") or "").strip()
        if not user_id or not name:
            raise DispatchValidationError(f"{role}必须是已明确指定姓名并完成平台映射的人员", status_code=409)
        recipients[user_id] = name
    if len(recipients) > MAX_OUTBOUND_RECIPIENTS:
        raise DispatchValidationError(
            f"单次最多只能向 {MAX_OUTBOUND_RECIPIENTS} 名明确指定人员发送消息",
            status_code=409,
        )
    return recipients


def enforce_outbound_audience_safety(
    feishu: Any,
    targets: list[tuple[str, str]],
    *,
    named_recipients: dict[str, str],
) -> None:
    """Fail closed unless the complete outbound audience is named and no larger than the hard limit."""
    audience: dict[str, str] = {}
    for receive_id, receive_id_type in dict.fromkeys(targets):
        normalized_id = str(receive_id or "").strip()
        if receive_id_type == "open_id":
            name = str(named_recipients.get(normalized_id) or "").strip()
            if not normalized_id or not name:
                raise DispatchValidationError("已拒绝向未明确指定姓名的人员发送消息", status_code=409)
            audience[normalized_id] = name
            continue
        if receive_id_type != "chat_id" or not normalized_id:
            raise DispatchValidationError("消息投递目标无效，已拒绝发送", status_code=409)
        try:
            directory = feishu.list_chat_members(normalized_id)
        except Exception as exc:
            raise DispatchValidationError("无法核验工作群成员，已按安全准则拒绝发送", status_code=409) from exc
        members = directory.get("members") if isinstance(directory, dict) else None
        if not isinstance(members, list):
            raise DispatchValidationError("工作群成员目录无效，已按安全准则拒绝发送", status_code=409)
        for member in members:
            member_id = str(member.get("member_id") or "").strip()
            member_name = str(member.get("name") or "").strip()
            if not member_id or not member_name or member_name == "未命名成员":
                raise DispatchValidationError("工作群存在未明确姓名的成员，已拒绝发送", status_code=409)
            audience[member_id] = member_name
        member_total = int(directory.get("member_total") or len(members))
        if member_total > MAX_OUTBOUND_RECIPIENTS:
            raise DispatchValidationError(
                f"工作群共有 {member_total} 人，超过单次 {MAX_OUTBOUND_RECIPIENTS} 人安全上限，已拒绝发送",
                status_code=409,
            )
    if len(audience) > MAX_OUTBOUND_RECIPIENTS:
        raise DispatchValidationError(
            f"本次消息将触达 {len(audience)} 人，超过单次 {MAX_OUTBOUND_RECIPIENTS} 人安全上限，已拒绝发送",
            status_code=409,
        )


def stored_delivery_policy(task: dict[str, Any]) -> dict[str, list[str]]:
    try:
        value = json.loads(str(task.get("delivery_policy_json") or "{}"))
    except (TypeError, ValueError):
        value = {}
    return normalize_delivery_policy(value, legacy_mode=str(task.get("delivery_mode") or "group"))


def public_dispatch_task(task: dict[str, Any]) -> dict[str, Any]:
    status = str(task.get("status") or "")
    claimed = bool(task.get("claimed_at")) or status in {"claimed", "pending_review", "returned", "completed"}
    reviewer_labels = {"waiting": "待编制", "pending": "待复核", "approved": "已通过", "rejected": "已退回"}
    participants = [{"role": "编制人", "name": str(task.get("assignee_name") or ""), "status": "已领取" if claimed else "待领取"}]
    participants.extend({
        "role": "复核人",
        "name": str(item.get("display_name") or ""),
        "status": reviewer_labels.get(str(item.get("status") or ""), str(item.get("status") or "")),
        "comment": str(item.get("comment") or ""),
    } for item in (task.get("_reviewers") or []))
    return {
        "task_id": str(task.get("task_id") or ""),
        "source_task_id": str(task.get("source_task_id") or ""),
        "task_name": str(task.get("task_name") or ""),
        "project_name": str(task.get("project_name") or ""),
        "skill": {
            "id": str(task.get("skill_id") or ""),
            "version": str(task.get("skill_version") or ""),
        },
        "delivery_mode": str(task.get("delivery_mode") or ""),
        "delivery_policy": stored_delivery_policy(task),
        "platform": str(task.get("platform_profile_id") or ""),
        "target_group_name": str(task.get("target_chat_name") or ""),
        "assignee_name": str(task.get("assignee_name") or ""),
        "participants": participants,
        "deadline": str(task.get("deadline") or ""),
        "status": str(task.get("status") or ""),
        "status_label": PUBLIC_STATUS_LABELS.get(str(task.get("status") or ""), str(task.get("status") or "")),
        "card_status": str(task.get("card_status") or "pending"),
        "file_status": str(task.get("file_status") or "pending"),
        "file_name": str(task.get("file_name") or ""),
        "template_version": str(task.get("template_version") or ""),
        "template_hash": str(task.get("template_hash") or "")[:12],
        "delivery_retry_count": int(task.get("delivery_retry_count") or 0),
        "error": sanitize_dispatch_error(task.get("delivery_error") or task.get("error") or ""),
        "created_at": str(task.get("created_at") or ""),
        "delivered_at": str(task.get("delivered_at") or ""),
        "claimed_at": str(task.get("claimed_at") or ""),
        "review_round": int(task.get("review_round") or 0),
        "review_card_status": str(task.get("review_card_status") or ""),
        "completion_card_status": str(task.get("completion_card_status") or ""),
        "completion_error": sanitize_dispatch_error(task.get("completion_error") or ""),
        "completed_at": str(task.get("completed_at") or ""),
        "submission_file_name": str(task.get("submission_file_name") or ""),
        "submission_delivery_status": str(task.get("submission_delivery_status") or ""),
        "can_retry": str(task.get("status") or "") == "dispatch_failed",
    }


def configured_platforms() -> list[dict[str, Any]]:
    return [
        {
            "profile_id": item["profile_id"],
            "label": item["label"],
            "domain_host": item.get("domain_host", ""),
            "configuration_ok": bool(item.get("configuration_ok")),
        }
        for item in feishu_app_bot.credential_profiles()
    ]


def build_service(
    *,
    registry: ProfessionalSkillRegistry,
    profile_id: str | None = None,
    store: ExternalDispatchStore | None = None,
    api_factory: Callable[..., Any] = feishu_app_bot.FeishuApi,
) -> ExternalTaskDispatchService:
    selected_profile = str(profile_id or feishu_app_bot.active_profile_id() or "").strip()
    credentials = feishu_app_bot.load_credentials(selected_profile)
    if not credentials.get("app_id") or not credentials.get("app_secret"):
        raise DispatchValidationError("当前投递平台机器人凭证未配置", status_code=409)
    issue = feishu_app_bot.credential_configuration_issue(selected_profile, credentials)
    if issue:
        raise DispatchValidationError(issue, status_code=409)
    defaults = feishu_app_bot.load_bot_defaults()
    verified_profiles = defaults.get("directDeliveryVerifiedProfiles") or []
    return ExternalTaskDispatchService(
        store=store or ExternalDispatchStore(),
        registry=registry,
        feishu=api_factory(
            credentials["app_id"],
            credentials["app_secret"],
            domain=credentials.get("domain") or feishu_app_bot.DEFAULT_FEISHU_DOMAIN,
        ),
        profile_id=selected_profile,
        authorized_group_name=str(defaults.get("authorizedDispatchGroupName") or AUTHORIZED_TEST_GROUP_NAME),
        app_url=feishu_app_bot.load_completion_card_app_url(),
        direct_delivery_verified=selected_profile in verified_profiles,
    )
